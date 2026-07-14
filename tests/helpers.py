from __future__ import annotations

import json
from io import BytesIO

import openpyxl


def workbook_bytes(*, comparison: bool = False) -> bytes:
    workbook = openpyxl.Workbook()
    default = workbook.active
    workbook.remove(default)

    pages = workbook.create_sheet("网页")
    queries = workbook.create_sheet("查询数")
    filters = workbook.create_sheet("过滤器")
    filters.append(["过滤器", "值"])
    filters.append(["日期", "过去 28 天"])

    if comparison:
        headers = [
            "维度",
            "过去 28 天 点击次数",
            "先前 28 天 点击次数",
            "过去 28 天 展示",
            "先前 28 天 展示",
            "过去 28 天 点击率",
            "先前 28 天 点击率",
            "过去 28 天 排名",
            "先前 28 天 排名",
        ]
        pages.append(headers)
        pages.append(
            ["https://laserpointerhub.com/blog/example", 2, 10, 120, 240, 0.0167, 0.0417, 12, 8]
        )
        pages.append(
            ["https://laserpointerhub.com/products/demo", 1, 1, 110, 100, 0.009, 0.01, 9, 10]
        )
        pages.append(
            ["https://laserpointerhub.com/blog/other", 5, 4, 180, 160, 0.0278, 0.025, 6, 7]
        )
        queries.append(headers)
        queries.append(["best example laser", 1, 1, 500, 350, 0.002, 0.0029, 14, 16])
    else:
        headers = ["维度", "点击次数", "展示", "点击率", "排名"]
        pages.append(headers)
        pages.append(["https://laserpointerhub.com/blog/example", 4, 100, 0.04, 9.5])
        queries.append(headers)
        queries.append(["example laser", 2, 50, 0.04, 10])

    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def blog_export_bytes() -> bytes:
    payload = {
        "type": "blogs",
        "exportedAt": "2026-07-14T00:00:00Z",
        "count": 2,
        "excludedFields": ["images"],
        "items": [
            {
                "id": "b1",
                "title": "Example Laser Guide",
                "slug": "example",
                "summary": "Example summary",
                "content": "Useful original content.",
                "tags": ["example"],
                "seoTitle": "Example Laser Guide",
                "seoDescription": "A useful example.",
                "seoKeywords": ["example laser"],
                "active": True,
                "createdAt": "2026-01-01T00:00:00Z",
                "updatedAt": "2026-07-01T00:00:00Z",
            },
            {
                "id": "b2",
                "title": "Other Laser Guide",
                "slug": "other",
                "summary": "Other summary",
                "content": "Other useful content.",
                "tags": ["other"],
                "seoTitle": "Other Laser Guide",
                "seoDescription": "Another example.",
                "seoKeywords": ["other laser"],
                "active": True,
                "createdAt": "2026-01-02T00:00:00Z",
                "updatedAt": "2026-07-02T00:00:00Z",
            },
        ],
    }
    return json.dumps(payload).encode()


def product_export_bytes() -> bytes:
    payload = {
        "type": "products",
        "exportedAt": "2026-07-14T00:00:00Z",
        "count": 1,
        "excludedFields": ["images"],
        "items": [
            {
                "id": "p1",
                "sku": "DEMO-1",
                "title": "Demo Laser",
                "slug": "demo",
                "price": 99,
                "inventory": 4,
                "attributes": {"wavelength": "520nm"},
                "description": "Demo product",
                "features": ["USB-C"],
                "packageList": ["Laser", "Cable"],
                "metaTitle": "Demo Laser",
                "metaDescription": "Demo product description",
                "keywords": ["demo laser"],
                "active": True,
                "createdAt": "2026-01-01T00:00:00Z",
                "updatedAt": "2026-07-01T00:00:00Z",
            }
        ],
    }
    return json.dumps(payload).encode()
