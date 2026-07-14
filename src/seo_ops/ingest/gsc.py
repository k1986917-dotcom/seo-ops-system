from __future__ import annotations

import warnings
from dataclasses import dataclass
from datetime import date, datetime
from io import BytesIO
from typing import Any

import openpyxl

SHEET_DIMENSIONS = {
    "图表": "date",
    "chart": "date",
    "查询数": "query",
    "queries": "query",
    "网页": "page",
    "pages": "page",
    "国家_地区": "country",
    "国家/地区": "country",
    "countries": "country",
    "设备": "device",
    "devices": "device",
    "搜索结果呈现": "search_appearance",
    "search appearance": "search_appearance",
}

FILTER_SHEETS = {"过滤器", "filters"}


@dataclass(frozen=True, slots=True)
class GSCMetric:
    dimension: str
    value: str
    period: str
    clicks: float | None
    impressions: float | None
    ctr: float | None
    position: float | None
    sheet_name: str
    source_row: int


@dataclass(frozen=True, slots=True)
class GSCParseResult:
    metrics: list[GSCMetric]
    metadata: dict[str, Any]


def _normal(value: Any) -> str:
    return str(value or "").strip().lower()


def _number(value: Any, *, percent: bool = False) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, str):
        clean = value.strip().replace(",", "")
        if not clean:
            return None
        if clean.endswith("%"):
            return float(clean[:-1]) / 100
        result = float(clean)
    else:
        result = float(value)
    if percent and result > 1:
        return result / 100
    return result


def _dimension_value(value: Any) -> str:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value or "").strip()


def _is_comparison(headers: tuple[Any, ...]) -> bool:
    text = " ".join(_normal(item) for item in headers)
    return len(headers) >= 9 or "先前" in text or "previous" in text


def _parse_dimension_sheet(sheet: Any, dimension: str) -> tuple[list[GSCMetric], bool]:
    row_iter = sheet.iter_rows(values_only=True)
    headers = tuple(next(row_iter, ()))
    comparison = _is_comparison(headers)
    metrics: list[GSCMetric] = []

    for row_number, row in enumerate(row_iter, start=2):
        if not row:
            continue
        value = _dimension_value(row[0] if len(row) > 0 else None)
        if not value:
            continue

        if comparison:
            padded = tuple(row) + (None,) * max(0, 9 - len(row))
            current = GSCMetric(
                dimension=dimension,
                value=value,
                period="current",
                clicks=_number(padded[1]),
                impressions=_number(padded[3]),
                ctr=_number(padded[5], percent=True),
                position=_number(padded[7]),
                sheet_name=sheet.title,
                source_row=row_number,
            )
            previous = GSCMetric(
                dimension=dimension,
                value=value,
                period="previous",
                clicks=_number(padded[2]),
                impressions=_number(padded[4]),
                ctr=_number(padded[6], percent=True),
                position=_number(padded[8]),
                sheet_name=sheet.title,
                source_row=row_number,
            )
            metrics.extend((current, previous))
        else:
            padded = tuple(row) + (None,) * max(0, 5 - len(row))
            metrics.append(
                GSCMetric(
                    dimension=dimension,
                    value=value,
                    period="current",
                    clicks=_number(padded[1]),
                    impressions=_number(padded[2]),
                    ctr=_number(padded[3], percent=True),
                    position=_number(padded[4]),
                    sheet_name=sheet.title,
                    source_row=row_number,
                )
            )
    return metrics, comparison


def parse_gsc_workbook(content: bytes) -> GSCParseResult:
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        workbook = openpyxl.load_workbook(BytesIO(content), read_only=True, data_only=True)

    metrics: list[GSCMetric] = []
    dimensions: list[str] = []
    filters: dict[str, str] = {}
    comparison = False

    try:
        for sheet_name in workbook.sheetnames:
            normalized = _normal(sheet_name)
            sheet = workbook[sheet_name]
            if normalized in FILTER_SHEETS:
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if row and row[0] not in (None, ""):
                        filters[str(row[0]).strip()] = str(row[1] if len(row) > 1 else "").strip()
                continue

            dimension = SHEET_DIMENSIONS.get(normalized)
            if not dimension:
                continue
            sheet_metrics, sheet_comparison = _parse_dimension_sheet(sheet, dimension)
            metrics.extend(sheet_metrics)
            dimensions.append(dimension)
            comparison = comparison or sheet_comparison
    finally:
        workbook.close()

    if not metrics:
        raise ValueError("工作簿中未找到可识别的 GSC 指标 Sheet")

    return GSCParseResult(
        metrics=metrics,
        metadata={
            "mode": "comparison" if comparison else "standard",
            "comparison": comparison,
            "dimensions": dimensions,
            "filters": filters,
            "metric_rows": len(metrics),
        },
    )
